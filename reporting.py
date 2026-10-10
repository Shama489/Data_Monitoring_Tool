"""Generate portable monitoring reports from persisted monitoring results."""

from __future__ import annotations

import io
import json
import re
import threading
from datetime import datetime, timezone
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

REPORT_FORMATS = {
    "pdf": {
        "mime_type": "application/pdf",
        "extension": "pdf",
    },
    "xlsx": {
        "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "extension": "xlsx",
    },
    "pptx": {
        "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "extension": "pptx",
    },
}
_chart_lock = threading.Lock()


def _safe_filename(value: Any) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value)).strip("._-")
    return (cleaned or "monitoring_report")[:100]


def _summary(result: dict[str, Any]) -> dict[str, Any]:
    monitoring = result.get("result", result)
    if not isinstance(monitoring, dict):
        raise ValueError("Monitoring result must be an object.")
    results = monitoring.get("results", {})
    if not isinstance(results, dict):
        results = {}
    root_cause = monitoring.get("root_cause", {})
    if not isinstance(root_cause, dict):
        root_cause = {}
    findings = root_cause.get("findings", [])
    if not isinstance(findings, list):
        findings = []
    notifications = monitoring.get("notifications", {})
    if not isinstance(notifications, dict):
        notifications = {}

    quality = results.get("quality", {})
    quality_report = quality.get("report", {}) if isinstance(quality, dict) else {}
    quality_score = quality.get("quality_score", {}) if isinstance(quality, dict) else {}
    drift = results.get("drift", {})
    anomalies = results.get("anomalies", {})
    return {
        "monitoring": monitoring,
        "results": results,
        "checks": monitoring.get("checks", []),
        "findings": [item for item in findings if isinstance(item, dict)],
        "quality_score": quality_score.get("overall_score") if isinstance(quality_score, dict) else None,
        "missing_values": quality_report.get("total_nulls") if isinstance(quality_report, dict) else None,
        "drift_score": drift.get("overall_drift_score") if isinstance(drift, dict) else None,
        "anomalies": anomalies.get("total_anomalies") if isinstance(anomalies, dict) else None,
        "notifications": notifications,
    }


def _chart_png(summary: dict[str, Any]) -> bytes:
    labels = ["Quality", "Drift", "Anomalies", "Findings"]
    values = [
        float(summary["quality_score"] or 0),
        float(summary["drift_score"] or 0),
        float(summary["anomalies"] or 0),
        float(len(summary["findings"])),
    ]
    with _chart_lock:
        figure, axis = plt.subplots(figsize=(8, 3.5))
        bars = axis.bar(labels, values, color=["#168f86", "#e0a03b", "#d95d5d", "#536d9e"])
        axis.set_title("Monitoring indicators")
        axis.set_ylabel("Score or count")
        axis.grid(axis="y", alpha=0.2)
        axis.bar_label(bars, fmt="%.1f", padding=3)
        figure.tight_layout()
        output = io.BytesIO()
        figure.savefig(output, format="png", dpi=150, bbox_inches="tight")
        plt.close(figure)
        return output.getvalue()


def _pdf_report(summary: dict[str, Any], title: str, chart: bytes) -> bytes:
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=title,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle",
        parent=styles["Title"],
        alignment=TA_LEFT,
        textColor=colors.HexColor("#126e66"),
        spaceAfter=12,
    ))
    story = [
        Paragraph(title, styles["ReportTitle"]),
        Paragraph(
            f"Generated {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Image(io.BytesIO(chart), width=7.0 * inch, height=3.05 * inch),
        Paragraph("Summary", styles["Heading2"]),
    ]
    summary_rows = [
        ["Monitoring checks", ", ".join(map(str, summary["checks"])) or "None"],
        ["Data quality score", _display(summary["quality_score"], "/100")],
        ["Missing values", _display(summary["missing_values"])],
        ["Drift score", _display(summary["drift_score"])],
        ["Anomalies", _display(summary["anomalies"])],
        ["Alert status", str(summary["notifications"].get("status", "not reported"))],
    ]
    table = Table(summary_rows, colWidths=[1.8 * inch, 5.2 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#e9f3f2")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#253238")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#c8d5d4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph("Findings", styles["Heading2"]))
    if summary["findings"]:
        for finding in summary["findings"]:
            signal = _escape(finding.get("signal", "finding"))
            evidence = _escape(json.dumps(finding.get("evidence", {}), ensure_ascii=False, default=str))
            recommendation = _escape(finding.get("recommended_investigation", ""))
            story.append(KeepTogether([
                Paragraph(f"<b>{signal.replace('_', ' ').title()}</b>", styles["BodyText"]),
                Paragraph(evidence, styles["Code"]),
                Paragraph(recommendation, styles["BodyText"]),
                Spacer(1, 6),
            ]))
    else:
        story.append(Paragraph("No findings were reported.", styles["BodyText"]))

    story.append(Paragraph("Alerts", styles["Heading2"]))
    alert_detail = _alert_summary(summary["notifications"])
    story.append(Paragraph(
        _escape(alert_detail),
        styles["BodyText"],
    ))
    story.append(PageBreak())
    story.append(Paragraph("Detailed monitoring results", styles["Heading2"]))
    story.append(Paragraph(
        _escape(json.dumps(summary["results"], ensure_ascii=False, indent=2, default=str)),
        styles["Code"],
    ))
    document.build(story)
    return output.getvalue()


def _excel_report(summary: dict[str, Any], title: str, chart: bytes) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_frame = pd.DataFrame([
            {"Metric": "Title", "Value": title},
            {"Metric": "Monitoring checks", "Value": ", ".join(map(str, summary["checks"]))},
            {"Metric": "Quality score", "Value": summary["quality_score"]},
            {"Metric": "Missing values", "Value": summary["missing_values"]},
            {"Metric": "Drift score", "Value": summary["drift_score"]},
            {"Metric": "Anomalies", "Value": summary["anomalies"]},
            {"Metric": "Alert status", "Value": summary["notifications"].get("status")},
            {"Metric": "Alert detail", "Value": summary["notifications"].get("error") or summary["notifications"].get("reason")},
        ])
        summary_frame.to_excel(writer, sheet_name="Summary", index=False)
        findings = summary["findings"]
        findings_frame = pd.DataFrame([
            {
                "Signal": finding.get("signal"),
                "Evidence": json.dumps(finding.get("evidence", {}), ensure_ascii=False, default=str),
                "Recommended investigation": finding.get("recommended_investigation"),
            }
            for finding in findings
        ], columns=["Signal", "Evidence", "Recommended investigation"])
        findings_frame.to_excel(writer, sheet_name="Findings", index=False)
        pd.DataFrame([summary["notifications"]]).to_excel(
            writer, sheet_name="Alerts", index=False
        )
        details = pd.json_normalize(summary["results"], sep=".").T.reset_index()
        details.columns = ["Metric", "Value"]
        details["Value"] = details["Value"].map(lambda value: json.dumps(value, default=str) if isinstance(value, (dict, list)) else value)
        details.to_excel(writer, sheet_name="Details", index=False)
        workbook = writer.book
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cells in sheet.columns:
                width = min(max(max(len(str(cell.value or "")) for cell in cells) + 2, 12), 72)
                sheet.column_dimensions[cells[0].column_letter].width = width
        from openpyxl.drawing.image import Image as ExcelImage

        chart_sheet = workbook.create_sheet("Charts")
        chart_sheet.add_image(ExcelImage(io.BytesIO(chart)), "A1")
    return output.getvalue()


def _pptx_report(summary: dict[str, Any], title: str, chart: bytes) -> bytes:
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    title_slide.placeholders[1].text = (
        f"Generated {datetime.now(timezone.utc).isoformat(timespec='seconds')}"
    )

    summary_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    summary_slide.shapes.title.text = "Monitoring summary"
    rows = [
        ("Checks", ", ".join(map(str, summary["checks"])) or "None"),
        ("Quality score", _display(summary["quality_score"], "/100")),
        ("Missing values", _display(summary["missing_values"])),
        ("Drift score", _display(summary["drift_score"])),
        ("Anomalies", _display(summary["anomalies"])),
        ("Alert status", str(summary["notifications"].get("status", "not reported"))),
    ]
    textbox = summary_slide.shapes.add_textbox(
        Inches(0.8), Inches(1.5), Inches(8.5), Inches(4.8)
    )
    frame = textbox.text_frame
    frame.clear()
    for index, (label, value) in enumerate(rows):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"{label}: {value}"
        paragraph.font.size = Pt(22)
        paragraph.space_after = Pt(13)

    chart_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    chart_slide.shapes.title.text = "Monitoring indicators"
    chart_slide.shapes.add_picture(io.BytesIO(chart), Inches(0.7), Inches(1.35), width=Inches(8.7))

    findings_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    findings_slide.shapes.title.text = "Findings and alerts"
    text = findings_slide.shapes.add_textbox(
        Inches(0.7), Inches(1.2), Inches(8.8), Inches(5.8)
    ).text_frame
    text.word_wrap = True
    entries = [
        f"{finding.get('signal', 'Finding')}: "
        f"{json.dumps(finding.get('evidence', {}), ensure_ascii=False, default=str)}"
        for finding in summary["findings"]
    ]
    alert = summary["notifications"]
    entries.append(
        _alert_summary(alert)
    )
    for index, entry in enumerate(entries or ["No findings were reported."]):
        paragraph = text.paragraphs[0] if index == 0 else text.add_paragraph()
        paragraph.text = entry
        paragraph.level = 0
        paragraph.font.size = Pt(14)
        paragraph.space_after = Pt(8)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _display(value: Any, suffix: str = "") -> str:
    return "Not available" if value is None else f"{value}{suffix}"


def _alert_summary(alert: dict[str, Any]) -> str:
    status = str(alert.get("status", "not reported"))
    sent = alert.get("sent", [])
    failed = alert.get("failed", [])
    sent_count = len(sent) if isinstance(sent, list) else 0
    failed_count = len(failed) if isinstance(failed, list) else 0
    details = alert.get("error") or alert.get("reason")
    summary = f"Status: {status}; {sent_count} delivered; {failed_count} failed"
    return f"{summary}; {details}" if details else summary


def _escape(value: str) -> str:
    from xml.sax.saxutils import escape

    return escape(value.encode("cp1252", errors="replace").decode("cp1252"))


def generate_monitoring_reports(
    monitoring_result: dict[str, Any],
    dataset_name: str,
    formats: list[str],
) -> dict[str, dict[str, Any]]:
    """Create reports in requested formats; raises on invalid formats or generation errors."""
    if not isinstance(monitoring_result, dict):
        raise ValueError("monitoring_result must be an object.")
    if not isinstance(formats, list) or not formats:
        raise ValueError("formats must be a non-empty list.")
    normalized_formats = list(dict.fromkeys(
        value.strip().lower() for value in formats if isinstance(value, str)
    ))
    unsupported = [value for value in normalized_formats if value not in REPORT_FORMATS]
    if unsupported:
        raise ValueError(f"Unsupported report formats: {unsupported}")
    if not normalized_formats:
        raise ValueError("formats must contain at least one supported format.")

    summary = _summary(monitoring_result)
    title = f"Monitoring Report — {_safe_filename(dataset_name)}"
    chart = _chart_png(summary)
    generators = {
        "pdf": _pdf_report,
        "xlsx": _excel_report,
        "pptx": _pptx_report,
    }
    reports = {}
    for format_name in normalized_formats:
        descriptor = REPORT_FORMATS[format_name]
        filename = f"{_safe_filename(dataset_name)}_monitoring_report.{descriptor['extension']}"
        reports[format_name] = {
            "content": generators[format_name](summary, title, chart),
            "filename": filename,
            "mime_type": descriptor["mime_type"],
        }
    return reports
